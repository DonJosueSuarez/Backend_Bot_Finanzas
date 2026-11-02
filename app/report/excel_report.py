from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.service.series import MESES, evolucion

AZUL_OSCURO = "1F3864"
AZUL = "2E75B6"
AZUL_CLARO = "DDEBF7"
GRIS = "7F7F7F"
GRIS_CLARO = "F2F2F2"
VERDE = "548235"
ROJO = "C00000"
NARANJA = "ED7D31"

MONEDA = '"$"#,##0.00'
PORCENTAJE = "0.0%"
FECHA = "dd/mm/yyyy"

NOMBRES_CATEGORIA = {
    "ALIMENTACION HOGAR": "Alimentación hogar",
    "PIQUEOS Y COMIDA FUERA": "Piqueos y comida fuera",
    "TRANSPORTE Y MOVILIDAD": "Transporte y movilidad",
    "SERVICIOS Y SUSCRIPCIONES": "Servicios y suscripciones",
    "CUIDADO PERSONAL": "Cuidado personal",
    "SALUD": "Salud",
    "ENTRETENIMIENTO Y SOCIAL": "Entretenimiento y social",
    "EDUCACION Y TRABAJO": "Educación y trabajo",
    "VARIOS E IMPREVISTOS": "Varios e imprevistos",
}

BORDE_FINO = Side(style="thin", color="D9D9D9")


def _fecha_larga(d: date) -> str:
    return f"{d.day} {MESES[d.month - 1]} {d.year}"


def _categoria(nombre) -> str:
    return NOMBRES_CATEGORIA.get(nombre, nombre) if nombre else "Sin categoría"


def _titulo(ws, titulo: str, subtitulo: str, ultima_columna: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_columna)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_columna)
    ws["A1"] = titulo
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=AZUL_OSCURO)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(size=10, italic=True, color=GRIS)
    ws["A2"].alignment = Alignment(indent=1)
    ws.row_dimensions[1].height = 34
    ws.sheet_view.showGridLines = False


def _encabezado(ws, fila: int, columna: int, textos: List[str]) -> None:
    for i, texto in enumerate(textos):
        celda = ws.cell(row=fila, column=columna + i, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center")


def _fila_tabla(ws, fila: int, valores: list, formatos: list) -> None:
    for i, (valor, formato) in enumerate(zip(valores, formatos), start=1):
        celda = ws.cell(row=fila, column=i, value=valor)
        celda.border = Border(bottom=BORDE_FINO)
        if formato:
            celda.number_format = formato


def _seccion(ws, fila: int, texto: str) -> None:
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(size=13, bold=True, color=AZUL_OSCURO)


def _sin_datos(ws, fila: int, texto: str) -> None:
    ws.cell(row=fila, column=1, value=texto).font = Font(italic=True, color=GRIS)


def _hoja_movimientos(ws, movimientos: List[Dict], subtitulo: str) -> None:
    ws.title = "Movimientos"
    ws.sheet_properties.tabColor = AZUL_OSCURO
    columnas = [("Fecha", 13), ("Tipo", 11), ("Descripción", 42), ("Categoría", 27), ("Necesidad / deseo", 18), ("Monto", 15)]
    _titulo(ws, "Reporte de movimientos", subtitulo, len(columnas))
    for i, (texto, ancho) in enumerate(columnas, start=1):
        ws.cell(row=4, column=i, value=texto)
        ws.column_dimensions[get_column_letter(i)].width = ancho

    if not movimientos:
        _encabezado(ws, 4, 1, [c[0] for c in columnas])
        _sin_datos(ws, 5, "No hay movimientos en este periodo.")
        return

    for fila, m in enumerate(movimientos, start=5):
        es_ingreso = m["tipo_movimiento"] == "INGRESO"
        valores = [
            m["fecha"],
            "Ingreso" if es_ingreso else "Egreso",
            m["descripcion"],
            _categoria(m["categoria"]) if not es_ingreso or m["categoria"] else "—",
            m["tipo_gasto"].capitalize() if m["tipo_gasto"] else "—",
            m["monto"],
        ]
        _fila_tabla(ws, fila, valores, [FECHA, None, None, None, None, MONEDA])
        ws.cell(row=fila, column=6).font = Font(color=VERDE if es_ingreso else ROJO)

    ultima = 4 + len(movimientos)
    tabla = Table(displayName="Movimientos", ref=f"A4:F{ultima}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = "A5"

    # Totales como valores (no fórmulas) para que se vean también en visores del teléfono
    ingresos = sum((m["monto"] for m in movimientos if m["tipo_movimiento"] == "INGRESO"), Decimal(0))
    egresos = sum((m["monto"] for m in movimientos if m["tipo_movimiento"] == "EGRESO"), Decimal(0))
    totales = [("Total ingresos", ingresos, VERDE), ("Total egresos", egresos, ROJO),
               ("Balance del periodo", ingresos - egresos, VERDE if ingresos >= egresos else ROJO)]
    for i, (texto, valor, color) in enumerate(totales):
        fila = ultima + 2 + i
        ws.cell(row=fila, column=5, value=texto).font = Font(bold=True)
        celda = ws.cell(row=fila, column=6, value=valor)
        celda.number_format = MONEDA
        celda.font = Font(bold=True, color=color)
        if i == 2:
            for col in (5, 6):
                ws.cell(row=fila, column=col).border = Border(top=Side(style="medium", color=AZUL_OSCURO))

    _ajustar_a_pagina(ws)
    ws.print_title_rows = "4:4"


def _tarjeta(ws, columna: int, etiqueta: str, valor, formato: str, color: str) -> None:
    for fila in (4, 5):
        ws.merge_cells(start_row=fila, start_column=columna, end_row=fila, end_column=columna + 1)
    arriba = ws.cell(row=4, column=columna, value=etiqueta)
    arriba.font = Font(size=10, bold=True, color="FFFFFF")
    arriba.fill = PatternFill("solid", fgColor=AZUL)
    arriba.alignment = Alignment(horizontal="center", vertical="center")
    abajo = ws.cell(row=5, column=columna, value=valor)
    abajo.font = Font(size=18, bold=True, color=color)
    abajo.fill = PatternFill("solid", fgColor=GRIS_CLARO)
    abajo.alignment = Alignment(horizontal="center", vertical="center")
    abajo.number_format = formato
    ws.cell(row=5, column=columna + 1).fill = PatternFill("solid", fgColor=GRIS_CLARO)


def _sin_ejes_ocultos(chart) -> None:
    # openpyxl 3.1 marca los ejes como borrados y Excel no los dibuja
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def _etiquetas(valor: bool = False, porcentaje: bool = False) -> DataLabelList:
    # Sin los flags explícitos, Excel agrega serie, categoría y leyenda a cada etiqueta
    return DataLabelList(showVal=valor, showPercent=porcentaje, showSerName=False,
                         showCatName=False, showLegendKey=False, showLeaderLines=False)


def _ajustar_a_pagina(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _hoja_dashboard(ws, movimientos: List[Dict], desde: date, hasta: date, subtitulo: str) -> None:
    ws.title = "Dashboard"
    ws.sheet_properties.tabColor = AZUL
    _titulo(ws, "Dashboard financiero", subtitulo, 12)
    ws.column_dimensions["A"].width = 26
    for col in range(2, 13):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 38

    egresos_lista = [m for m in movimientos if m["tipo_movimiento"] == "EGRESO"]
    ingresos = sum((m["monto"] for m in movimientos if m["tipo_movimiento"] == "INGRESO"), Decimal(0))
    egresos = sum((m["monto"] for m in egresos_lista), Decimal(0))
    dias = (hasta - desde).days + 1

    _tarjeta(ws, 1, "Ingresos", ingresos, MONEDA, VERDE)
    _tarjeta(ws, 3, "Egresos", egresos, MONEDA, ROJO)
    _tarjeta(ws, 5, "Balance", ingresos - egresos, MONEDA, VERDE if ingresos >= egresos else ROJO)
    _tarjeta(ws, 7, "Movimientos", len(movimientos), "0", AZUL_OSCURO)
    _tarjeta(ws, 9, "Gasto diario promedio", egresos / dias, MONEDA, AZUL_OSCURO)
    ws.column_dimensions["I"].width = 15
    ws.column_dimensions["J"].width = 15

    # --- Gastos por categoría
    fila = 8
    _seccion(ws, fila, "Gastos por categoría")
    por_categoria = defaultdict(Decimal)
    for m in egresos_lista:
        por_categoria[_categoria(m["categoria"])] += m["monto"]
    categorias = sorted(por_categoria.items(), key=lambda c: c[1], reverse=True)

    _encabezado(ws, fila + 1, 1, ["Categoría", "Total", "% del gasto"])
    if categorias:
        for i, (nombre, total) in enumerate(categorias):
            _fila_tabla(ws, fila + 2 + i, [nombre, total, float(total / egresos)], [None, MONEDA, PORCENTAJE])
        ultima = fila + 1 + len(categorias)
        ws.conditional_formatting.add(f"B{fila + 2}:B{ultima}", DataBarRule(start_type="min", end_type="max", color=AZUL))

        barras = BarChart()
        barras.type = "bar"
        barras.title = "Gasto por categoría"
        barras.legend = None
        barras.x_axis.scaling.orientation = "maxMin"  # la categoría más alta arriba
        barras.add_data(Reference(ws, min_col=2, min_row=fila + 1, max_row=ultima), titles_from_data=True)
        barras.set_categories(Reference(ws, min_col=1, min_row=fila + 2, max_row=ultima))
        barras.series[0].graphicalProperties.solidFill = AZUL
        barras.series[0].graphicalProperties.line.solidFill = AZUL
        barras.dataLabels = _etiquetas(valor=True)
        barras.width, barras.height = 17, 7.5
        _sin_ejes_ocultos(barras)
        # Los montos ya van en las etiquetas: el eje de valores sobra
        barras.y_axis.delete = True
        barras.y_axis.majorGridlines = None
        ws.add_chart(barras, "E8")
    else:
        _sin_datos(ws, fila + 2, "Sin gastos en el periodo.")

    # --- Necesidad vs deseo
    fila = 25
    _seccion(ws, fila, "Necesidad vs deseo")
    _encabezado(ws, fila + 1, 1, ["Tipo de gasto", "Total", "% del gasto"])
    tipos = [(t.capitalize(), sum((m["monto"] for m in egresos_lista if m["tipo_gasto"] == t), Decimal(0)))
             for t in ("NECESIDAD", "DESEO")]
    if egresos:
        for i, (nombre, total) in enumerate(tipos):
            _fila_tabla(ws, fila + 2 + i, [nombre, total, float(total / egresos)], [None, MONEDA, PORCENTAJE])

        dona = DoughnutChart()
        dona.title = "Necesidad vs deseo"
        dona.holeSize = 55
        dona.add_data(Reference(ws, min_col=2, min_row=fila + 1, max_row=fila + 3), titles_from_data=True)
        dona.set_categories(Reference(ws, min_col=1, min_row=fila + 2, max_row=fila + 3))
        for idx, color in enumerate((AZUL, NARANJA)):
            punto = DataPoint(idx=idx)
            punto.graphicalProperties.solidFill = color
            dona.series[0].dPt.append(punto)
        dona.dataLabels = _etiquetas(porcentaje=True)
        dona.legend.position = "b"
        dona.width, dona.height = 11, 7.5
        ws.add_chart(dona, "E25")
    else:
        _sin_datos(ws, fila + 2, "Sin gastos en el periodo.")

    # --- Evolución en el tiempo
    fila = 42
    serie = evolucion(movimientos, desde, hasta)
    titulo = "Evolución diaria" if serie["agrupacion"] == "dia" else "Evolución mensual"
    puntos = [(p["etiqueta"], p["ingresos"], p["egresos"]) for p in serie["puntos"]]
    _seccion(ws, fila, titulo)
    _encabezado(ws, fila + 1, 1, ["Periodo", "Ingresos", "Egresos"])
    for i, punto in enumerate(puntos):
        _fila_tabla(ws, fila + 2 + i, list(punto), [None, MONEDA, MONEDA])
    ultima = fila + 1 + len(puntos)

    if not movimientos:
        _ajustar_a_pagina(ws)
        return

    # Sin título: la sección ya lo tiene y la leyenda va arriba
    linea = LineChart()
    linea.y_axis.numFmt = MONEDA
    linea.legend.position = "t"
    linea.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9"))
    linea.add_data(Reference(ws, min_col=2, max_col=3, min_row=fila + 1, max_row=ultima), titles_from_data=True)
    linea.set_categories(Reference(ws, min_col=1, min_row=fila + 2, max_row=ultima))
    for serie, color in zip(linea.series, (VERDE, ROJO)):
        serie.graphicalProperties.line.solidFill = color
        serie.graphicalProperties.line.width = 28000
        serie.smooth = False
    linea.width, linea.height = 17, 7.5
    _sin_ejes_ocultos(linea)
    ws.add_chart(linea, "E42")
    _ajustar_a_pagina(ws)


def build_excel_report(movimientos: List[Dict], desde: date, hasta: date, generado: datetime) -> bytes:
    """Hoja 1: movimientos del periodo. Hoja 2: dashboard con indicadores y gráficos."""
    subtitulo = f"Periodo: {_fecha_larga(desde)} – {_fecha_larga(hasta)}   ·   Generado: {generado:%d/%m/%Y %H:%M}"

    wb = Workbook()
    _hoja_movimientos(wb.active, movimientos, subtitulo)
    _hoja_dashboard(wb.create_sheet(), movimientos, desde, hasta, subtitulo)
    wb.properties.title = "Reporte de movimientos"
    wb.properties.creator = "Bot de finanzas"

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
